from docx import Document
from PIL import Image

from eventfile.report import PhotoEvidence, ReportData, build_report
from eventfile.vision import AttendanceRow, EventDetails


def img(path, color="gray"):
    Image.new("RGB", (800, 600), color).save(path)
    return path


def sample(tmp_path, attendees=()):
    return ReportData(
        institution="Sample College of Engineering",
        details=EventDetails(
            title="Workshop on Building with Open-Source AI",
            organiser="Coding Club, Department of CSE",
            date="2026-10-14",
            time="10:00 AM - 1:00 PM",
            venue="Seminar Hall 2",
            resource_person="Dr. R. Meenakshi",
            resource_person_role="Senior Research Engineer, Chennai AI Labs",
            event_type="Workshop",
        ),
        participants=22,
        objective="To introduce students to open-weight models.",
        summary="Students installed Ollama and ran Gemma.",
        outcome="Three teams built a chatbot.",
        photos=[
            PhotoEvidence(img(tmp_path / "p1.jpg"), "Inauguration — 14 Oct 2026, 10:08 AM"),
            PhotoEvidence(img(tmp_path / "p2.jpg"), "Hands-on session — 14 Oct 2026, 11:26 AM"),
        ],
        poster=img(tmp_path / "poster.jpg", "white"),
        attendance_sheets=[img(tmp_path / "a1.jpg", "beige")],
        attendees=list(attendees),
    )


def all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_report_has_details_narrative_and_captions(tmp_path):
    out = build_report(sample(tmp_path), tmp_path / "report.docx")
    text = all_text(Document(out))
    for expected in [
        "SAMPLE COLLEGE OF ENGINEERING",
        "Workshop on Building with Open-Source AI",
        "14 October 2026",
        "Dr. R. Meenakshi, Senior Research Engineer, Chennai AI Labs",
        "22",
        "To introduce students to open-weight models.",
        "Inauguration — 14 Oct 2026, 10:08 AM",
        "Event Coordinator",
    ]:
        assert expected in text


def test_report_embeds_photos_poster_and_sheets(tmp_path):
    doc = Document(build_report(sample(tmp_path), tmp_path / "report.docx"))
    assert len(doc.inline_shapes) == 2 + 1 + 1


def test_attendee_list_only_when_given(tmp_path):
    without = all_text(Document(build_report(sample(tmp_path), tmp_path / "a.docx")))
    assert "List of Participants" not in without

    people = [AttendanceRow(1, "Priya S", "21CS001", True), AttendanceRow(2, "Arun K", "", False)]
    doc = Document(build_report(sample(tmp_path, people), tmp_path / "b.docx"))
    text = all_text(doc)
    assert "List of Participants" in text and "Priya S" in text and "21CS001" in text
