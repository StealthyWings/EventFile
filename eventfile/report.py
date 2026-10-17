"""Write the event report as a Word file, in the layout most college IQAC offices ask for."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor

from .vision import AttendanceRow, EventDetails


@dataclass
class PhotoEvidence:
    image: Path
    caption: str


@dataclass
class ReportData:
    details: EventDetails
    participants: int
    objective: str
    summary: str
    outcome: str
    photos: list[PhotoEvidence]
    institution: str = ""
    poster: Path | None = None
    attendance_sheets: list[Path] = field(default_factory=list)
    attendees: list[AttendanceRow] = field(default_factory=list)


def _date_text(details: EventDetails) -> str:
    d = details.event_date
    return f"{d.day} {d:%B %Y}" if d else details.date


def _heading(doc, text, size=13):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    p.paragraph_format.space_before = Pt(12)
    return p


def _centered(doc, text, size, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return p


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _picture(doc, path, width_in, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width_in))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = c.add_run(caption)
        run.italic = True
        run.font.size = Pt(10)


def build_report(data: ReportData, out_path) -> Path:
    out_path = Path(out_path)
    d = data.details
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.core_properties.title = f"Event report: {d.title}"
    doc.core_properties.comments = "Drafted with EventFile (Gemma 4, run locally). Check before signing."

    if data.institution:
        _centered(doc, data.institution.upper(), 15, bold=True)
    if d.organiser:
        _centered(doc, d.organiser, 12)
    _centered(doc, "EVENT REPORT", 14, bold=True)

    rows = [
        ("Name of the Event", d.title),
        ("Type of Event", d.event_type),
        ("Organised by", d.organiser),
        ("Date", _date_text(d)),
        ("Time", d.time),
        ("Venue", d.venue),
        ("Resource Person", ", ".join(x for x in (d.resource_person, d.resource_person_role) if x)),
        ("Number of Participants", str(data.participants)),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value or "—"
        cells[0].paragraphs[0].runs[0].bold = True
        cells[0].width, cells[1].width = Inches(2.2), Inches(4.3)

    for title, body in (
        ("Objective", data.objective),
        ("Summary of the Event", data.summary),
        ("Outcome", data.outcome),
    ):
        _heading(doc, title)
        doc.add_paragraph(body or "—")

    if data.photos:
        _heading(doc, "Photographs (geotagged)")
        for i, photo in enumerate(data.photos, 1):
            _picture(doc, photo.image, 5.0, f"Figure {i}: {photo.caption}")

    _heading(doc, "Signatures")
    sig = doc.add_table(rows=2, cols=3)
    for i, role in enumerate(("Event Coordinator", "Head of the Department", "Principal")):
        sig.cell(0, i).text = "\n\n\n"
        sig.cell(1, i).text = role
        sig.cell(1, i).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    annexure = 1
    if data.poster:
        _page_break(doc)
        _heading(doc, f"Annexure {annexure}: Poster / Brochure")
        _picture(doc, data.poster, 5.0)
        annexure += 1

    if data.attendance_sheets:
        _page_break(doc)
        _heading(doc, f"Annexure {annexure}: Attendance Sheets")
        for sheet in data.attendance_sheets:
            _picture(doc, sheet, 5.5)
        annexure += 1

    if data.attendees:
        _page_break(doc)
        _heading(doc, f"Annexure {annexure}: List of Participants")
        people = doc.add_table(rows=1, cols=3)
        people.style = "Table Grid"
        for cell, text in zip(people.rows[0].cells, ("S.No", "Name", "Register Number")):
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
        for i, person in enumerate(data.attendees, 1):
            cells = people.add_row().cells
            cells[0].text, cells[1].text, cells[2].text = str(i), person.name, person.register_number

    doc.save(out_path)
    return out_path
